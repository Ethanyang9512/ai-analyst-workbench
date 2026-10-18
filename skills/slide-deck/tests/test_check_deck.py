# -*- coding: utf-8 -*-
"""check_deck.py 的冒烟测试（需要 python-pptx；没装就整组跳过）。

运行（在技能目录下）：${AAW_PY:-python3} tests/test_check_deck.py -v
片子在临时目录里现造，不依赖任何母版；文字和数字都是虚构的。
"""
import os
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # 技能目录（本文件在 tests/ 下）
SCRIPT = os.path.join(HERE, "scripts", "check_deck.py")

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Cm, Pt
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


@unittest.skipUnless(HAS_PPTX, "没装 python-pptx，check_deck 测试跳过")
class TestCheckDeck(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="check_deck_test_")

    def tearDown(self):
        self.tmp.cleanup()

    def deck(self, name, bar_color="1A1A1A", box_top_cm=3.0, box_text="续费率 W6 起比前 5 周低约 4pt"):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Cm(33.87), Cm(19.05)
        slide = prs.slides.add_slide(prs.slide_layouts[6])            # 空白版式
        tb = slide.shapes.add_textbox(Cm(2), Cm(box_top_cm), Cm(20), Cm(1.2))
        tb.text_frame.text = box_text
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        data = CategoryChartData()
        data.categories = ["W5", "W6", "W7"]
        data.add_series("续费率", (80.3, 76.4, 76.1))
        gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Cm(2), Cm(6), Cm(16), Cm(8), data)
        fill = gf.chart.plots[0].series[0].format.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(bar_color)
        path = os.path.join(self.tmp.name, name)
        prs.save(path)
        return path

    def run_cd(self, *args):
        r = subprocess.run([sys.executable, SCRIPT] + list(args), capture_output=True, text=True, encoding="utf-8")
        return r.returncode, r.stdout + r.stderr

    def test_clean_deck_passes(self):
        code, out = self.run_cd(self.deck("ok.pptx"))
        self.assertEqual(code, 0, out)
        self.assertIn("全部通过", out)
        self.assertIn("页脚条 未给，不查", out)

    def test_footer_only_checked_when_given(self):
        path = self.deck("low.pptx", box_top_cm=17.6)                  # 框底沿 18.8cm，在页面内
        self.assertEqual(self.run_cd(path)[0], 0)
        code, out = self.run_cd(path, "--footer-top", "18.16")
        self.assertEqual(code, 1, out)
        self.assertIn("压页脚条", out)

    def test_off_palette_chart_color_is_hint_only(self):
        path = self.deck("brand.pptx", bar_color="0055AA")
        code, out = self.run_cd(path)
        self.assertEqual(code, 0, out)                                 # 只提示，不计入失败
        self.assertIn("#0055AA", out)
        code, out = self.run_cd(path, "--palette", "0055AA,FFFFFF")
        self.assertNotIn("#0055AA", out)
        code, out = self.run_cd(path, "--palette", "none")
        self.assertNotIn("配色表以外", out)

    def test_terms_blacklist(self):
        path = self.deck("terms.pptx", box_text="这里有一个自造词：抓手")
        code, out = self.run_cd(path, "--terms", "抓手,赋能")
        self.assertEqual(code, 1, out)
        self.assertIn("抓手", out)


if __name__ == "__main__":
    unittest.main()
